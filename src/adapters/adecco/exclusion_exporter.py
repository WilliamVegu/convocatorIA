"""Official 5-column Excel exporter for Adecco portfolio and exclusion reports (Ley N° 29733)."""
from __future__ import annotations

import io
from typing import List, Optional, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from src.ports.contracts import CarteraExclusionRow5Col


class ExclusionReportExporter:
    """Exports standardized 5-column Excel reports guaranteed to exclude all private contact/financial data."""

    def export_to_excel(
        self,
        rows: List[CarteraExclusionRow5Col | dict],
        cuenta_cliente: Optional[str] = None,
    ) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        title_suffix = f" - {cuenta_cliente}" if cuenta_cliente else " - Consolidado"
        ws.title = f"Exclusiones{title_suffix}"[:31]

        # Ensure gridlines visible
        ws.views.sheetView[0].showGridLines = True

        # Headers - Exactly 5 columns
        headers = [
            "DNI",
            "Nombres y Apellidos",
            "Perfil",
            "Vigencia de Exclusión",
            "Estado",
        ]
        ws.append(headers)

        # Style header row
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="0A192F", end_color="0A192F", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin", color="CCCCCC"),
            right=Side(style="thin", color="CCCCCC"),
            top=Side(style="thin", color="CCCCCC"),
            bottom=Side(style="thin", color="CCCCCC"),
        )
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")

        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

        ws.row_dimensions[1].height = 28

        # Populate data rows
        data_font = Font(name="Segoe UI", size=10)
        alt_fill = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")

        for r_idx, row_item in enumerate(rows, 2):
            if isinstance(row_item, CarteraExclusionRow5Col):
                d = row_item.model_dump()
            else:
                d = row_item

            values = [
                d.get("dni", ""),
                d.get("nombres_y_apellidos", ""),
                d.get("perfil", ""),
                d.get("vigencia_exclusion", ""),
                d.get("estado", ""),
            ]
            ws.append(values)
            ws.row_dimensions[r_idx].height = 22

            for c_idx in range(1, 6):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.font = data_font
                cell.border = thin_border
                if r_idx % 2 == 1:
                    cell.fill = alt_fill
                if c_idx in {1, 4, 5}:
                    cell.alignment = center_align
                else:
                    cell.alignment = left_align

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()
