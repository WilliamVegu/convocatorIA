import pptx
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
import os

TEMPLATE_PATH = r"C:\Users\willi\Downloads\PPT_EJEMPL.pptx"
OUT_PATH_DOWNLOADS = r"C:\Users\willi\Downloads\ATS_Sistema_Reclutamiento_IA.pptx"
OUT_PATH_LOCAL = r"c:\Users\willi\Downloads\ats\ATS_Sistema_Reclutamiento_IA.pptx"

IMG_FICHA = r"c:\Users\willi\Downloads\ats\assets\screen_ficha.png"
IMG_ADECCO = r"c:\Users\willi\Downloads\ats\assets\screen_adecco.png"

prs = pptx.Presentation(TEMPLATE_PATH)

def set_p_text(p, text, font_name="Century Gothic", size=None, bold=None, color=None, align=None):
    if p.runs:
        r = p.runs[0]
        r.text = text
        for extra in p.runs[1:]:
            extra.text = ""
    else:
        r = p.add_run()
        r.text = text
        
    if font_name:
        r.font.name = font_name
    if size is not None:
        r.font.size = Pt(size)
    if bold is not None:
        r.font.bold = bold
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    if align is not None:
        p.alignment = align

def clean_paragraph_indent(p):
    pPr = p._p.get_or_add_pPr()
    for attr in ["marL", "indent", "lvl"]:
        if attr in pPr.attrib:
            del pPr.attrib[attr]
    for child in list(pPr):
        if "bu" in child.tag:
            pPr.remove(child)
    buNone = OxmlElement("a:buNone")
    pPr.append(buNone)

def update_shape_paragraphs(shape, items):
    tf = shape.text_frame
    for i, item in enumerate(items):
        if i < len(tf.paragraphs):
            p = tf.paragraphs[i]
        else:
            p = tf.add_paragraph()
        set_p_text(
            p,
            text=item.get("text", ""),
            font_name=item.get("font", "Century Gothic"),
            size=item.get("size", None),
            bold=item.get("bold", None),
            color=item.get("color", None),
            align=item.get("align", None)
        )
    for extra_p in tf.paragraphs[len(items):]:
        extra_p.text = ""

def find_shape(slide, name, shape_id=None, min_left=None):
    for s in slide.shapes:
        if shape_id is not None and s.shape_id == shape_id:
            return s
        if min_left is not None and s.name == name and s.left / 914400 >= min_left:
            return s
        if shape_id is None and min_left is None and s.name == name:
            return s
    return None

# ==================== SLIDE 1 ====================
s1 = prs.slides[0]

sh = find_shape(s1, "Text 3")
if sh: update_shape_paragraphs(sh, [{"text": "AI LAB PERÚ · TATA CONSULTANCY SERVICES", "size": 10.3, "bold": True, "color": "00D7FF"}])

sh = find_shape(s1, "Text 4")
if sh:
    sh.width = Inches(7.2)
    update_shape_paragraphs(sh, [
        {"text": "ATS Core System", "size": 48.0, "bold": True, "color": "F7FAFC"},
        {"text": "Reclutamiento con IA", "size": 48.0, "bold": True, "color": "F7FAFC"}
    ])

sh = find_shape(s1, "Text 5")
if sh: update_shape_paragraphs(sh, [{"text": "De 35h semanales en Excel a un ATS Hexagonal con IA responsable, deduplicación y control de proveedores.", "size": 15.5, "color": "F7FAFC"}])

sh = find_shape(s1, "Text 16")
if sh: update_shape_paragraphs(sh, [{"text": "RETO TÉCNICO", "size": 10.0, "bold": True, "color": "FF8A3D"}])

sh = find_shape(s1, "Text 17")
if sh: update_shape_paragraphs(sh, [{"text": "Negocio + IA + Evidencia", "size": 22.0, "bold": True, "color": "F7FAFC"}])

sh = find_shape(s1, "Text 20")
if sh: update_shape_paragraphs(sh, [{"text": "Detectar", "size": 6.8, "color": "AEB8C7"}])

sh = find_shape(s1, "Text 23")
if sh: update_shape_paragraphs(sh, [{"text": "Extraer", "size": 6.8, "color": "AEB8C7"}])

sh = find_shape(s1, "Text 26")
if sh: update_shape_paragraphs(sh, [{"text": "Validar", "size": 6.8, "color": "AEB8C7"}])

sh = find_shape(s1, "Text 28")
if sh: update_shape_paragraphs(sh, [{"text": "Decidir", "size": 6.8, "color": "AEB8C7"}])

sh = find_shape(s1, "Text 29")
if sh: update_shape_paragraphs(sh, [{"text": "La propuesta no busca “hacer IA por moda”. Busca erradicar fricción operativa, ordenar datos propios y blindar control humano con trazabilidad.", "size": 11.5, "color": "F7FAFC"}])

sh = find_shape(s1, "Text 31")
if sh: update_shape_paragraphs(sh, [{"text": "AI Lab Perú · TCS Reclutamiento IA", "size": 8.5, "color": "AEB8C7"}])

# ==================== SLIDE 2 ====================
s2 = prs.slides[1]

sh = find_shape(s2, "Text 4")
if sh: update_shape_paragraphs(sh, [{"text": "Dolor y oportunidad", "size": 32.0, "bold": True, "color": "F7FAFC"}])

sh = find_shape(s2, "Text 10")
if sh: update_shape_paragraphs(sh, [{"text": "35 horas/semana en Excel: 25h en registro manual de candidatos, 5h actualizando estados y 5h en cruce manual con Adecco.", "size": 10.5, "bold": True, "color": "00D7FF"}])

sh = find_shape(s2, "Text 15")
if sh: update_shape_paragraphs(sh, [{"text": "Proveedor a ciegas (10% útil): El 90% de CVs de Adecco se descarta por candidatos ya vistos, evaluados o filtros internos.", "size": 10.5, "bold": True, "color": "34D399"}])

sh = find_shape(s2, "Text 20")
if sh: update_shape_paragraphs(sh, [{"text": "Datos fragmentados y errores críticos: BD de 11 pestañas con errores #DIV/0!, candidatos con 127 años y duplicados continuos.", "size": 10.5, "bold": True, "color": "9B5CFF"}])

sh = find_shape(s2, "Text 25")
if sh: update_shape_paragraphs(sh, [{"text": "Paradoja de IA en LinkedIn: Hiring Assistant contratado rinde +51% en respuesta, pero solo se usa en 7% por falta de configuración.", "size": 10.5, "bold": True, "color": "FF8A3D"}])

sh = find_shape(s2, "Text 27")
if sh: update_shape_paragraphs(sh, [{"text": "Tesis para el jurado: No automatizamos por automatizar; construimos un ATS propio para gobernar datos, erradicar retrabajo y proteger cumplimiento legal (Ley 29733).", "size": 11.2, "color": "F7FAFC"}])

sh = find_shape(s2, "Text 28")
if sh: update_shape_paragraphs(sh, [{"text": "AI Lab Perú · TCS Reclutamiento IA", "size": 8.5, "color": "AEB8C7"}])

# ==================== SLIDE 3 ====================
s3 = prs.slides[2]

sh = find_shape(s3, "Text 3")
if sh: update_shape_paragraphs(sh, [{"text": "MÉTODO DE TRABAJO", "size": 10.5, "bold": True, "color": "00D7FF"}])

sh = find_shape(s3, "Text 4")
if sh: update_shape_paragraphs(sh, [{"text": "De la idea al impacto", "size": 29.0, "bold": True, "color": "F7FAFC"}])

sh = find_shape(s3, "Text 5")
if sh: update_shape_paragraphs(sh, [{"text": "Mantenemos el esquema del reto: cuatro fases centradas en valor, usuario y evidencia.", "size": 12.2, "color": "AEB8C7"}])

# Phase 1
sh = find_shape(s3, "Text 7");  sh and update_shape_paragraphs(sh, [{"text": "1", "size": 18.0, "bold": True, "color": "2F80FF"}])
sh = find_shape(s3, "Text 8");  sh and update_shape_paragraphs(sh, [{"text": "DISCOVERY", "size": 9.6, "bold": True, "color": "2F80FF"}])
sh = find_shape(s3, "Text 10"); sh and update_shape_paragraphs(sh, [{"text": "Entender el problema", "size": 14.1, "bold": True, "color": "F7FAFC"}])
sh = find_shape(s3, "Text 11"); sh and update_shape_paragraphs(sh, [{"text": "Levantamiento con recruiters, auditoría de licencias LinkedIn, causa raíz del 10% y mapeo de 35h en Excel.", "size": 9.5, "color": "AEB8C7"}])
sh = find_shape(s3, "Text 12"); sh and update_shape_paragraphs(sh, [{"text": "Entregables: Fichas de proceso AS IS, actas de discovery y catálogo de 12 capacidades funcionales.", "size": 9.0, "color": "F7FAFC"}])

# Phase 2
sh = find_shape(s3, "Text 15"); sh and update_shape_paragraphs(sh, [{"text": "2", "size": 18.0, "bold": True, "color": "34D399"}])
sh = find_shape(s3, "Text 16"); sh and update_shape_paragraphs(sh, [{"text": "IDEACIÓN", "size": 9.6, "bold": True, "color": "34D399"}])
sh = find_shape(s3, "Text 18"); sh and update_shape_paragraphs(sh, [{"text": "Diseñar la solución", "size": 14.1, "bold": True, "color": "F7FAFC"}])
sh = find_shape(s3, "Text 19"); sh and update_shape_paragraphs(sh, [{"text": "Spec-Driven Development (GitHub Spec Kit), Arquitectura Hexagonal, DDL relacional (10 tablas) y contratos Pydantic v2.", "size": 9.5, "color": "AEB8C7"}])
sh = find_shape(s3, "Text 20"); sh and update_shape_paragraphs(sh, [{"text": "Entregables: spec.md, plan.md, data-model.md, research.md y contratos de interfaz tipados.", "size": 9.0, "color": "F7FAFC"}])

# Phase 3
sh = find_shape(s3, "Text 23"); sh and update_shape_paragraphs(sh, [{"text": "3", "size": 18.0, "bold": True, "color": "9B5CFF"}])
sh = find_shape(s3, "Text 24"); sh and update_shape_paragraphs(sh, [{"text": "BUILDING", "size": 9.6, "bold": True, "color": "9B5CFF"}])
sh = find_shape(s3, "Text 26"); sh and update_shape_paragraphs(sh, [{"text": "Construir con IA", "size": 14.1, "bold": True, "color": "F7FAFC"}])
sh = find_shape(s3, "Text 27"); sh and update_shape_paragraphs(sh, [{"text": "Core Python 3.12, SQLite WAL, adaptadores IA (Gemini/Grok/Offline), UI Streamlit y 131 tests unitarios.", "size": 9.5, "color": "AEB8C7"}])
sh = find_shape(s3, "Text 28"); sh and update_shape_paragraphs(sh, [{"text": "Entregables: ATS funcional, motor de deduplicación, simulador CTC y consola de auditoría inmutable.", "size": 9.0, "color": "F7FAFC"}])

# Phase 4
sh = find_shape(s3, "Text 31"); sh and update_shape_paragraphs(sh, [{"text": "4", "size": 18.0, "bold": True, "color": "FF8A3D"}])
sh = find_shape(s3, "Text 32"); sh and update_shape_paragraphs(sh, [{"text": "DEMO", "size": 9.6, "bold": True, "color": "FF8A3D"}])
sh = find_shape(s3, "Text 34"); sh and update_shape_paragraphs(sh, [{"text": "Mostrar y aprender", "size": 14.1, "bold": True, "color": "F7FAFC"}])
sh = find_shape(s3, "Text 35"); sh and update_shape_paragraphs(sh, [{"text": "Flujo E2E validado: Ficha DNI, parsing CV, semáforo Adecco, reporte de exclusión y bitácora forense.", "size": 9.5, "color": "AEB8C7"}])
sh = find_shape(s3, "Text 36"); sh and update_shape_paragraphs(sh, [{"text": "Entregables: App productiva en vivo, 131 tests al 100%, guía quickstart y plan de escala corporativo.", "size": 9.0, "color": "F7FAFC"}])

# Bottom principles
sh = find_shape(s3, "Text 37"); sh and update_shape_paragraphs(sh, [{"text": "Principios que elevaron la calidad", "size": 12.5, "bold": True, "color": "F7FAFC"}])
sh = find_shape(s3, "Text 38"); sh and update_shape_paragraphs(sh, [{"text": "Usuario", "size": 9.3, "bold": True, "color": "00D7FF"}])
sh = find_shape(s3, "Text 40"); sh and update_shape_paragraphs(sh, [{"text": "Datos", "size": 9.3, "bold": True, "color": "34D399"}])
sh = find_shape(s3, "Text 42"); sh and update_shape_paragraphs(sh, [{"text": "IA responsable", "size": 9.3, "bold": True, "color": "9B5CFF"}])
sh = find_shape(s3, "Text 44"); sh and update_shape_paragraphs(sh, [{"text": "Evidencia", "size": 9.3, "bold": True, "color": "FF8A3D"}])

sh = find_shape(s3, "Text 46"); sh and update_shape_paragraphs(sh, [{"text": "AI Lab Perú · TCS Reclutamiento IA", "size": 8.5, "color": "AEB8C7"}])

# ==================== SLIDE 4 ====================
s4 = prs.slides[3]

sh = find_shape(s4, "Text 3"); sh and update_shape_paragraphs(sh, [{"text": "CONSTRUCCIÓN DE LA SOLUCIÓN", "size": 10.5, "bold": True, "color": "00D7FF"}])
sh = find_shape(s4, "Text 4"); sh and update_shape_paragraphs(sh, [{"text": "IA guiada por especificaciones", "size": 29.0, "bold": True, "color": "F7FAFC"}])
sh = find_shape(s4, "Text 5"); sh and update_shape_paragraphs(sh, [{"text": "El valor no vino de una sola herramienta: vino de orquestar personas, specs, modelos y validación.", "size": 12.2, "color": "AEB8C7"}])

# Left card - Rigorous and true to the project
sh = find_shape(s4, "Text 7"); sh and update_shape_paragraphs(sh, [{"text": "Spec Driven Development", "size": 16.0, "bold": True, "color": "F7FAFC"}])
sh = find_shape(s4, "Text 8"); sh and update_shape_paragraphs(sh, [{"text": "•", "size": 12.5, "bold": True, "color": "00D7FF"}])
sh = find_shape(s4, "Text 9", min_left=1.0); sh and update_shape_paragraphs(sh, [{"text": "Specify / Spec Kit: Formalización de la Constitución, historias de usuario y especificaciones ejecutables (spec.md, plan.md).", "size": 10.0, "color": "F7FAFC"}])
sh = find_shape(s4, "Text 10"); sh and update_shape_paragraphs(sh, [{"text": "•", "size": 12.5, "bold": True, "color": "34D399"}])
sh = find_shape(s4, "Text 11"); sh and update_shape_paragraphs(sh, [{"text": "Claude Code / Antigravity: Ensamblado de Arquitectura Hexagonal en capas, DDL SQLAlchemy y validadores Pydantic v2.", "size": 10.0, "color": "F7FAFC"}])
sh = find_shape(s4, "Text 12"); sh and update_shape_paragraphs(sh, [{"text": "•", "size": 12.5, "bold": True, "color": "FF8A3D"}])
sh = find_shape(s4, "Text 13"); sh and update_shape_paragraphs(sh, [{"text": "Google AI Studio / Gemini: Calibración y benchmarking de prompts para extracción estructurada de CVs con Gemini 2.5 Flash.", "size": 10.0, "color": "F7FAFC"}])
sh = find_shape(s4, "Text 14"); sh and update_shape_paragraphs(sh, [{"text": "•", "size": 12.5, "bold": True, "color": "9B5CFF"}])
sh = find_shape(s4, "Text 15"); sh and update_shape_paragraphs(sh, [{"text": "Pytest Automation: Suite integral de 131 tests automáticos (unitarios, de contrato e integración) pasando al 100%.", "size": 10.0, "color": "F7FAFC"}])

sh = find_shape(s4, "Text 16"); sh and update_shape_paragraphs(sh, [{"text": "Modelo y runtime usados", "size": 11.2, "bold": True, "color": "00D7FF"}])
sh = find_shape(s4, "Text 17"); sh and update_shape_paragraphs(sh, [{"text": "Google Gemini 2.5 Flash · xAI Grok-2 · Python 3.12 · Streamlit · SQLAlchemy 2.0 · SQLite WAL · Pydantic v2 · RapidFuzz", "size": 9.2, "color": "AEB8C7"}])

# Right card - Real features, clean layout, top anchor, no weird indents
sh = find_shape(s4, "Text 19"); sh and update_shape_paragraphs(sh, [{"text": "La Solución: ATS Core MVP", "size": 16.0, "bold": True, "color": "F7FAFC"}])

right_box = find_shape(s4, "Text 9", shape_id=54)
if not right_box:
    right_box = find_shape(s4, "Text 9", min_left=6.0)

if right_box:
    right_box.width = Inches(4.30)
    right_box.top = Inches(2.62)
    tf = right_box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.TOP
    
    features = [
        ("Capacidades del sistema web en producción:", 12.5, True, "F7FAFC"),
        ("• Ficha Única: Autollenado DNI (<5ms), normalización E.164 (+519XXXXXXXX) y canal WhatsApp Web 1-clic.", 9.8, False, "F7FAFC"),
        ("• Extracción CV/CUL: Multi-modelo IA (Gemini/Grok) y fallback offline con censura ética (Ley 29733).", 9.8, False, "F7FAFC"),
        ("• Validador masivo Adecco: Ingesta tolerante a alias, semáforo (Rojo/Amarillo/Verde) y alerta Boomerang.", 9.8, False, "F7FAFC"),
        ("• Reporte oficial de exclusiones: Generación en 1-clic de 5 columnas estandarizadas y anonimizadas.", 9.8, False, "F7FAFC"),
        ("• Simulador Financiero CTC: Factor legal 1.56 (D.L. 728), 21% neto-bruto y guardas con 0.00% #DIV/0!.", 9.8, False, "F7FAFC"),
        ("• Screening telefónico 7D: Registro humano estructurado de 7 dimensiones y alertas de conmutación.", 9.8, False, "F7FAFC"),
        ("• Seguridad & Auditoría: Control RBAC de 4 roles corporativos y bitácora inmutable append-only.", 9.8, False, "F7FAFC"),
    ]
    
    # Remove all paragraphs beyond what we need
    while len(tf.paragraphs) > len(features):
        p_extra = tf.paragraphs[-1]
        p_extra._p.getparent().remove(p_extra._p)
        
    for i, (f_text, f_size, f_bold, f_color) in enumerate(features):
        if i < len(tf.paragraphs):
            p = tf.paragraphs[i]
        else:
            p = tf.add_paragraph()
        clean_paragraph_indent(p)
        p.space_before = Pt(4)
        p.space_after = Pt(2)
        set_p_text(p, text=f_text, font_name="Century Gothic", size=f_size, bold=f_bold, color=f_color)

# Update badges on the right with clean typography
for s in s4.shapes:
    if s.name == "Grupo 40":
        for sub in s.shapes:
            if sub.has_text_frame and sub.name == "Text 21":
                update_shape_paragraphs(sub, [{"text": "Ficha DNI &\nWhatsApp 1-Clic", "size": 7.5, "color": "F7FAFC"}])
    elif s.name == "Grupo 37":
        for sub in s.shapes:
            if sub.has_text_frame and sub.name == "Text 24":
                update_shape_paragraphs(sub, [{"text": "Extracción CV\nMulti-Modelo", "size": 7.5, "color": "F7FAFC"}])
    elif s.name == "Grupo 39":
        for sub in s.shapes:
            if sub.has_text_frame and sub.name == "Text 27":
                update_shape_paragraphs(sub, [{"text": "Validador Masivo\nPlanillas Adecco", "size": 7.5, "color": "F7FAFC"}])
    elif s.name == "Grupo 38":
        for sub in s.shapes:
            if sub.has_text_frame and sub.name == "Text 30":
                update_shape_paragraphs(sub, [{"text": "Simulador CTC &\nAuditoría Forense", "size": 7.5, "color": "F7FAFC"}])

sh = find_shape(s4, "Text 32"); sh and update_shape_paragraphs(sh, [{"text": "AI Lab Perú · TCS Reclutamiento IA", "size": 8.5, "color": "AEB8C7"}])

# ==================== SLIDE 5 ====================
s5 = prs.slides[4]

sh = find_shape(s5, "Text 3"); sh and update_shape_paragraphs(sh, [{"text": "CIERRE GANADOR", "size": 10.5, "bold": True, "color": "00D7FF"}])
sh = find_shape(s5, "Text 4"); sh and update_shape_paragraphs(sh, [{"text": "Evidencia, impacto y escalabilidad", "size": 29.0, "bold": True, "color": "F7FAFC"}])
sh = find_shape(s5, "Text 5"); sh and update_shape_paragraphs(sh, [{"text": "Una solución técnica real convertida en activo estratégico para Reclutamiento TCS.", "size": 12.2, "color": "AEB8C7"}])

# Replace images
for s in s5.shapes:
    if s.name == "Imagen 36" and os.path.exists(IMG_FICHA):
        rId = s._element.xpath('.//a:blip')[0].attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed']
        img_part = s5.part.related_part(rId)
        with open(IMG_FICHA, "rb") as f_img:
            img_part._blob = f_img.read()
    elif s.name == "Imagen 40" and os.path.exists(IMG_ADECCO):
        rId = s._element.xpath('.//a:blip')[0].attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed']
        img_part = s5.part.related_part(rId)
        with open(IMG_ADECCO, "rb") as f_img:
            img_part._blob = f_img.read()

sh = find_shape(s5, "Text 17"); sh and update_shape_paragraphs(sh, [{"text": "Highlights", "size": 18.0, "bold": True, "color": "F7FAFC"}])

# Highlights 3 cards
sh = find_shape(s5, "Text 19"); sh and update_shape_paragraphs(sh, [{"text": "Impacto de negocio", "size": 13.0, "bold": True, "color": "00D7FF"}])
sh = find_shape(s5, "Text 20"); sh and update_shape_paragraphs(sh, [{"text": "Erradica 35h/sem de carga manual en Excel, eleva la tasa útil de CVs de Adecco de 10% a >60% y evita comisiones indebidas con alerta Boomerang.", "size": 9.3, "color": "F7FAFC"}])

sh = find_shape(s5, "Text 22"); sh and update_shape_paragraphs(sh, [{"text": "Solidez técnica", "size": 13.0, "bold": True, "color": "9B5CFF"}])
sh = find_shape(s5, "Text 23"); sh and update_shape_paragraphs(sh, [{"text": "Arquitectura Hexagonal desacoplada, 131 tests automáticos pasando (100%), SQLite WAL / PostgreSQL ready y resiliencia offline total.", "size": 9.3, "color": "F7FAFC"}])

sh = find_shape(s5, "Text 25"); sh and update_shape_paragraphs(sh, [{"text": "Escala responsable", "size": 13.0, "bold": True, "color": "34D399"}])
sh = find_shape(s5, "Text 26"); sh and update_shape_paragraphs(sh, [{"text": "Cumplimiento estricto Ley 29733 (datos personales censurados), supervisión humana innegociable (HITL), cero scraping y auditoría inmutable.", "size": 9.3, "color": "F7FAFC"}])

sh = find_shape(s5, "Text 28"); sh and update_shape_paragraphs(sh, [{"text": "Siguiente paso ejecutivo: Desplegar piloto en cuenta real de TCS Perú, medir horas recuperadas y viabilizar integración oficial RSC.", "size": 10.5, "bold": True, "color": "F7FAFC"}])

sh = find_shape(s5, "Text 29"); sh and update_shape_paragraphs(sh, [{"text": "AI Lab Perú · TCS Reclutamiento IA", "size": 8.5, "color": "AEB8C7"}])

prs.save(OUT_PATH_DOWNLOADS)
prs.save(OUT_PATH_LOCAL)
print("Successfully generated refined presentations in:")
print(f"  -> {OUT_PATH_DOWNLOADS}")
print(f"  -> {OUT_PATH_LOCAL}")
