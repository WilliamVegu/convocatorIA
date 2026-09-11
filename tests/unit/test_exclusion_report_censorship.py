"""Unit tests for Ley 29733 censorship and official 5-column structure of the exclusion report."""
import io
import pytest
import openpyxl

from src.adapters.adecco.exclusion_exporter import ExclusionReportExporter
from src.ports.contracts import CarteraExclusionRow5Col


def test_exporter_exact_5_columns():
    rows = [
        CarteraExclusionRow5Col(
            dni="76128709",
            nombres_y_apellidos="Diego Alonso Ramos Quispe",
            perfil="Senior Java Backend Developer",
            vigencia_exclusion="Hasta 15/12/2026",
            estado="En Proceso Activo",
        ),
        CarteraExclusionRow5Col(
            dni="46753314",
            nombres_y_apellidos="Carlos Eduardo Garcia Sanchez",
            perfil="DevOps Specialist",
            vigencia_exclusion="Hasta 30/11/2026",
            estado="Descarte Temporal",
        ),
    ]

    exporter = ExclusionReportExporter()
    excel_bytes = exporter.export_to_excel(rows, cuenta_cliente="BCP")

    # Load workbook from bytes
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    ws = wb.active

    headers = [cell.value for cell in ws[1]]
    assert len(headers) == 5
    assert headers == ["DNI", "Nombres y Apellidos", "Perfil", "Vigencia de Exclusión", "Estado"]

    # Verify rows count
    assert ws.max_row == 3  # 1 header + 2 data rows


def test_exporter_absolute_censorship_under_ley_29733():
    rows = [
        CarteraExclusionRow5Col(
            dni="76128709",
            nombres_y_apellidos="Diego Alonso Ramos Quispe",
            perfil="Senior Java Backend Developer",
            vigencia_exclusion="Hasta 15/12/2026",
            estado="En Proceso Activo",
        )
    ]

    exporter = ExclusionReportExporter()
    excel_bytes = exporter.export_to_excel(rows)

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    ws = wb.active

    all_cell_values = []
    for row in ws.iter_rows(values_only=True):
        for val in row:
            if val is not None:
                all_cell_values.append(str(val).lower())

    joined_text = " ".join(all_cell_values)
    # Check that contact details and financial fields are completely absent
    assert "@" not in joined_text
    assert "+51" not in joined_text
    assert "sueldo" not in joined_text
    assert "salario" not in joined_text
    assert "tarifa" not in joined_text
    assert "ctc" not in joined_text
